import io
from pathlib import Path
from typing import Optional
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from backend.models.schemas import MeetingMinutes, MeetingDetail


class ExporterService:
    @staticmethod
    def to_markdown(meeting: MeetingDetail) -> str:
        """Retorna a ata completa em Markdown."""
        if not meeting.summary:
            return "# Ata não disponível."

        content = [meeting.summary.raw_markdown]
        content.append("\n\n---\n## 🎙️ Transcrição Completa dos Diálogos\n")
        for seg in meeting.segments:
            m = int(seg.start // 60)
            s = int(seg.start % 60)
            time_str = f"{m:02d}:{s:02d}"
            content.append(f"**[{time_str}] {seg.speaker}:** {seg.text}  ")

        return "\n".join(content)

    @staticmethod
    def to_plain_text(meeting: MeetingDetail) -> str:
        """Gera versão em texto simples."""
        if not meeting.summary:
            return "Ata não disponível."

        s = meeting.summary
        lines = [
            f"ATA DE REUNIÃO: {s.title.upper()}",
            f"Data: {s.date}",
            f"Duração: {s.duration_minutes:.1f} minutos",
            f"Participantes: {', '.join(s.participants)}",
            "=" * 60,
            "\n1. RESUMO EXECUTIVO:",
            s.executive_summary,
            "\n2. TÓPICOS DISCUTIDOS:",
        ]

        for i, t in enumerate(s.main_topics, 1):
            lines.append(f"  {i}. {t.title}")
            lines.append(f"     Discussão: {t.discussion}")
            if t.conclusions:
                lines.append(f"     Conclusão: {t.conclusions}")

        lines.append("\n3. DECISÕES TOMADAS:")
        for d in s.decisions:
            lines.append(f"  - {d}")

        lines.append("\n4. PLANO DE AÇÃO:")
        for a in s.action_items:
            lines.append(f"  - [ ] {a.task} | Responsável: {a.owner} | Prazo: {a.deadline} | Status: {a.status}")

        if s.open_points:
            lines.append("\n5. PONTOS EM ABERTO:")
            for p in s.open_points:
                lines.append(f"  - {p}")

        lines.append("\n" + "=" * 60)
        lines.append("TRANSCRIÇÃO DETALHADA:")
        for seg in meeting.segments:
            m = int(seg.start // 60)
            sec = int(seg.start % 60)
            lines.append(f"[{m:02d}:{sec:02d}] {seg.speaker}: {seg.text}")

        return "\n".join(lines)

    @staticmethod
    def to_docx_bytes(meeting: MeetingDetail) -> io.BytesIO:
        """Gera um arquivo DOCX (Microsoft Word) estilizado profissionalmente."""
        doc = Document()

        # Configurações de margem
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

        summary = meeting.summary

        # Título
        title_p = doc.add_paragraph()
        title_run = title_p.add_run(summary.title if summary else meeting.title)
        title_run.font.name = "Arial"
        title_run.font.size = Pt(20)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(30, 41, 59)
        title_p.paragraph_format.space_after = Pt(4)

        # Subtítulo com metadados
        sub_p = doc.add_paragraph()
        sub_run = sub_p.add_run(f"Data: {summary.date if summary else meeting.created_at} | Duração: {meeting.audio_duration / 60:.1f} min")
        sub_run.font.size = Pt(10)
        sub_run.font.italic = True
        sub_run.font.color.rgb = RGBColor(100, 116, 139)
        sub_p.paragraph_format.space_after = Pt(14)

        if summary:
            # Participantes
            p_part = doc.add_paragraph()
            p_part.add_run("Participantes: ").bold = True
            p_part.add_run(", ".join(summary.participants))
            p_part.paragraph_format.space_after = Pt(12)

            # Resumo Executivo
            h1 = doc.add_heading("1. Resumo Executivo", level=1)
            h1.paragraph_format.space_before = Pt(12)
            p_exec = doc.add_paragraph(summary.executive_summary)
            p_exec.paragraph_format.space_after = Pt(12)

            # Tópicos Principais
            doc.add_heading("2. Tópicos Discutidos", level=1)
            for idx, topic in enumerate(summary.main_topics, 1):
                p_topic = doc.add_paragraph()
                p_topic.add_run(f"{idx}. {topic.title}").bold = True
                doc.add_paragraph(f"Discussão: {topic.discussion}")
                if topic.conclusions:
                    doc.add_paragraph(f"Conclusão: {topic.conclusions}")

            # Decisões Tomadas
            doc.add_heading("3. Decisões Tomadas", level=1)
            for dec in summary.decisions:
                p_dec = doc.add_paragraph(style="List Bullet")
                p_dec.add_run(dec)

            # Matriz de Ações (Tabela no Word)
            doc.add_heading("4. Plano de Ação", level=1)
            if summary.action_items:
                table = doc.add_table(rows=1, cols=4)
                table.style = "Table Grid"
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                hdr_cells = table.rows[0].cells
                hdr_titles = ["Tarefa / Ação", "Responsável", "Prazo", "Status"]
                for i, title in enumerate(hdr_titles):
                    hdr_cells[i].text = title
                    for paragraph in hdr_cells[i].paragraphs:
                        for run in paragraph.runs:
                            run.font.bold = True
                            run.font.size = Pt(10)

                for item in summary.action_items:
                    row_cells = table.add_row().cells
                    row_cells[0].text = item.task
                    row_cells[1].text = item.owner
                    row_cells[2].text = item.deadline
                    row_cells[3].text = item.status

            # Pontos em Aberto
            if summary.open_points:
                doc.add_heading("5. Pontos em Aberto", level=1)
                for pt in summary.open_points:
                    doc.add_paragraph(pt, style="List Bullet")

        # Transcrição Completa
        doc.add_page_break()
        doc.add_heading("Transcrição Integral da Reunião", level=1)
        for seg in meeting.segments:
            m = int(seg.start // 60)
            s = int(seg.start % 60)
            p = doc.add_paragraph()
            time_run = p.add_run(f"[{m:02d}:{s:02d}] ")
            time_run.font.color.rgb = RGBColor(148, 163, 184)
            time_run.font.size = Pt(9)

            speaker_run = p.add_run(f"{seg.speaker}: ")
            speaker_run.bold = True

            text_run = p.add_run(seg.text)

        target_stream = io.BytesIO()
        doc.save(target_stream)
        target_stream.seek(0)
        return target_stream

    @staticmethod
    def _format_srt_time(seconds: float) -> str:
        """Converte segundos para o formato de tempo do SRT (HH:MM:SS,mmm)."""
        hrs = int(seconds // 3600)
        mins = int((seconds % 3600) // 60)
        secs = int(seconds % 60)
        millis = int(round((seconds - int(seconds)) * 1000))
        if millis >= 1000:
            millis = 999
        return f"{hrs:02d}:{mins:02d}:{secs:02d},{millis:03d}"

    @staticmethod
    def transcript_to_markdown(meeting: MeetingDetail) -> str:
        """Retorna a transcrição dos diálogos em formato Markdown."""
        title = meeting.summary.title if meeting.summary else meeting.title
        date_str = meeting.summary.date if meeting.summary else meeting.created_at
        duration_min = meeting.audio_duration / 60.0
        unique_speakers = sorted(list(set(seg.speaker for seg in meeting.segments)))

        content = [
            f"# Transcrição: {title}",
            "",
            f"**Data:** {date_str}  ",
            f"**Duração:** {duration_min:.1f} minutos  ",
            f"**Locutores:** {', '.join(unique_speakers)}  ",
            "",
            "---",
            "",
            "## 🎙️ Diálogos Integrais",
            ""
        ]

        for seg in meeting.segments:
            m = int(seg.start // 60)
            s = int(seg.start % 60)
            time_str = f"{m:02d}:{s:02d}"
            content.append(f"**[{time_str}] {seg.speaker}:** {seg.text}  ")

        return "\n".join(content)

    @staticmethod
    def transcript_to_plain_text(meeting: MeetingDetail) -> str:
        """Gera versão da transcrição em texto simples (.txt)."""
        title = meeting.summary.title if meeting.summary else meeting.title
        date_str = meeting.summary.date if meeting.summary else meeting.created_at
        duration_min = meeting.audio_duration / 60.0
        unique_speakers = sorted(list(set(seg.speaker for seg in meeting.segments)))

        lines = [
            f"TRANSCRIÇÃO DE REUNIÃO: {title.upper()}",
            f"Data: {date_str}",
            f"Duração: {duration_min:.1f} minutos",
            f"Locutores: {', '.join(unique_speakers)}",
            "=" * 60,
            ""
        ]

        for seg in meeting.segments:
            m = int(seg.start // 60)
            s = int(seg.start % 60)
            lines.append(f"[{m:02d}:{s:02d}] {seg.speaker}: {seg.text}")

        return "\n".join(lines)

    @staticmethod
    def transcript_to_docx_bytes(meeting: MeetingDetail) -> io.BytesIO:
        """Gera um arquivo DOCX elegante contendo a transcrição integral dos diálogos."""
        doc = Document()

        for section in doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

        title = meeting.summary.title if meeting.summary else meeting.title
        date_str = meeting.summary.date if meeting.summary else meeting.created_at
        duration_min = meeting.audio_duration / 60.0
        unique_speakers = sorted(list(set(seg.speaker for seg in meeting.segments)))

        # Título
        title_p = doc.add_paragraph()
        title_run = title_p.add_run(f"Transcrição: {title}")
        title_run.font.name = "Arial"
        title_run.font.size = Pt(20)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(30, 41, 59)
        title_p.paragraph_format.space_after = Pt(4)

        # Subtítulo com metadados
        sub_p = doc.add_paragraph()
        sub_run = sub_p.add_run(
            f"Data: {date_str} | Duração: {duration_min:.1f} min | Locutores: {', '.join(unique_speakers)}"
        )
        sub_run.font.name = "Arial"
        sub_run.font.size = Pt(10)
        sub_run.font.italic = True
        sub_run.font.color.rgb = RGBColor(100, 116, 139)
        sub_p.paragraph_format.space_after = Pt(16)

        # Heading Diálogos
        h1 = doc.add_heading("Diálogos da Reunião", level=1)
        h1.paragraph_format.space_before = Pt(8)
        h1.paragraph_format.space_after = Pt(12)

        for seg in meeting.segments:
            m = int(seg.start // 60)
            s = int(seg.start % 60)
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15

            time_run = p.add_run(f"[{m:02d}:{s:02d}] ")
            time_run.font.name = "Arial"
            time_run.font.size = Pt(9.5)
            time_run.font.color.rgb = RGBColor(148, 163, 184)

            speaker_run = p.add_run(f"{seg.speaker}: ")
            speaker_run.font.name = "Arial"
            speaker_run.font.size = Pt(10.5)
            speaker_run.bold = True
            speaker_run.font.color.rgb = RGBColor(37, 99, 235)

            text_run = p.add_run(seg.text)
            text_run.font.name = "Arial"
            text_run.font.size = Pt(10.5)
            text_run.font.color.rgb = RGBColor(30, 41, 59)

        target_stream = io.BytesIO()
        doc.save(target_stream)
        target_stream.seek(0)
        return target_stream

    @classmethod
    def transcript_to_srt(cls, meeting: MeetingDetail) -> str:
        """Gera legendas sincronizadas no formato universal SRT."""
        srt_blocks = []
        for idx, seg in enumerate(meeting.segments, 1):
            start_str = cls._format_srt_time(seg.start)
            end_str = cls._format_srt_time(seg.end)
            text = f"[{seg.speaker}] {seg.text}"
            srt_blocks.append(f"{idx}\n{start_str} --> {end_str}\n{text}\n")
        return "\n".join(srt_blocks)

